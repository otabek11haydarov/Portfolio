import os
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_resume_docx(data, filename):
    doc = docx.Document()
    
    # Set page margins to 0.5 inches (Harvard standard for 1-page fit)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Base font setup
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(10)
    font.color.rgb = RGBColor(0, 0, 0)

    # Helper: Add bottom border to heading paragraph
    def add_bottom_border(paragraph):
        pPr = paragraph._p.get_or_add_pPr()
        pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                         r'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="000000"/>'
                         r'</w:pBdr>')
        pPr.append(pBdr)

    # Helper: Section Header
    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.name = 'Times New Roman'
        add_bottom_border(p)
        return p

    # Helper: 2-column header row (Company/University on left, City on right)
    def add_entry_header(left_title, right_info, left_sub="", right_sub=""):
        table = doc.add_table(rows=0, cols=2)
        table.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        
        # Row 1
        row1 = table.add_row()
        cell_l1 = row1.cells[0]
        cell_r1 = row1.cells[1]
        cell_l1.width = Inches(5.6)
        cell_r1.width = Inches(1.9)
        
        p_l1 = cell_l1.paragraphs[0]
        p_l1.paragraph_format.space_before = Pt(3)
        p_l1.paragraph_format.space_after = Pt(0)
        run_l1 = p_l1.add_run(left_title)
        run_l1.bold = True
        run_l1.font.size = Pt(10)
        run_l1.font.name = 'Times New Roman'
        
        p_r1 = cell_r1.paragraphs[0]
        p_r1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_r1.paragraph_format.space_before = Pt(3)
        p_r1.paragraph_format.space_after = Pt(0)
        run_r1 = p_r1.add_run(right_info)
        run_r1.bold = True
        run_r1.font.size = Pt(10)
        run_r1.font.name = 'Times New Roman'
        
        # Row 2 (if provided)
        if left_sub or right_sub:
            row2 = table.add_row()
            cell_l2 = row2.cells[0]
            cell_r2 = row2.cells[1]
            cell_l2.width = Inches(5.6)
            cell_r2.width = Inches(1.9)
            
            p_l2 = cell_l2.paragraphs[0]
            p_l2.paragraph_format.space_before = Pt(0)
            p_l2.paragraph_format.space_after = Pt(1)
            run_l2 = p_l2.add_run(left_sub)
            run_l2.italic = True
            run_l2.font.size = Pt(9.5)
            run_l2.font.name = 'Times New Roman'
            
            p_r2 = cell_r2.paragraphs[0]
            p_r2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_r2.paragraph_format.space_before = Pt(0)
            p_r2.paragraph_format.space_after = Pt(1)
            run_r2 = p_r2.add_run(right_sub)
            run_r2.italic = True
            run_r2.font.size = Pt(9.5)
            run_r2.font.name = 'Times New Roman'

    # Helper: Bullet point
    def add_bullet(text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.left_indent = Inches(0.2)
        run = p.add_run(text)
        run.font.size = Pt(9)
        run.font.name = 'Times New Roman'

    # --- Header (Name & Contact) ---
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(1)
    run_name = p_name.add_run(data['name'].upper())
    run_name.bold = True
    run_name.font.size = Pt(16)
    run_name.font.name = 'Times New Roman'

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.space_after = Pt(4)
    run_contact = p_contact.add_run(data['contact_line'])
    run_contact.font.size = Pt(9)
    run_contact.font.name = 'Times New Roman'

    # --- Education ---
    add_section_header(data['edu_title'])
    for edu in data['education']:
        add_entry_header(edu['institution'], edu['location'], edu['degree'], edu['dates'])
        for bullet in edu['bullets']:
            add_bullet(bullet)

    # --- Experience / Projects ---
    add_section_header(data['exp_title'])
    for exp in data['experience']:
        add_entry_header(exp['company'], exp['location'], exp['role'], exp['dates'])
        for bullet in exp['bullets']:
            add_bullet(bullet)

    # --- Leadership ---
    add_section_header(data['lead_title'])
    for lead in data['leadership']:
        add_entry_header(lead['company'], lead['location'], lead['role'], lead['dates'])
        for bullet in lead['bullets']:
            add_bullet(bullet)

    # --- Skills & Interests ---
    add_section_header(data['skills_title'])
    for s_line in data['skills_lines']:
        p_skill = doc.add_paragraph()
        p_skill.paragraph_format.space_before = Pt(0)
        p_skill.paragraph_format.space_after = Pt(1)
        p_skill.paragraph_format.line_spacing = 1.05
        run_bold = p_skill.add_run(s_line['label'] + ": ")
        run_bold.bold = True
        run_bold.font.size = Pt(9)
        run_bold.font.name = 'Times New Roman'
        run_val = p_skill.add_run(s_line['value'])
        run_val.font.size = Pt(9)
        run_val.font.name = 'Times New Roman'

    doc.save(filename)
    print(f"Generated DOCX: {filename}")

def create_resume_html(data, filename):
    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<style>
  @page {{
    size: letter;
    margin: 0.45in 0.5in;
  }}
  * {{
    box-sizing: border-box;
    margin: 0;
    padding: 0;
  }}
  body {{
    font-family: 'Times New Roman', Times, serif;
    font-size: 9.5pt;
    line-height: 1.25;
    color: #111;
    background-color: #fff;
  }}
  .header {{
    text-align: center;
    margin-bottom: 6px;
  }}
  .name {{
    font-size: 16pt;
    font-weight: bold;
    letter-spacing: 0.5px;
    margin-bottom: 2px;
    text-transform: uppercase;
  }}
  .contact {{
    font-size: 8.8pt;
    color: #222;
  }}
  .contact a {{
    color: inherit;
    text-decoration: none;
  }}
  .section-title {{
    font-size: 10.5pt;
    font-weight: bold;
    text-transform: uppercase;
    border-bottom: 1px solid #111;
    margin-top: 7px;
    margin-bottom: 3px;
    padding-bottom: 1px;
    letter-spacing: 0.4px;
  }}
  .item {{
    margin-bottom: 3px;
  }}
  .item-header {{
    display: flex;
    justify-content: space-between;
    font-size: 9.5pt;
    font-weight: bold;
  }}
  .item-sub {{
    display: flex;
    justify-content: space-between;
    font-size: 9pt;
    font-style: italic;
    color: #222;
    margin-bottom: 1px;
  }}
  ul {{
    margin: 1px 0 2px 18px;
    padding-left: 0;
  }}
  li {{
    font-size: 8.8pt;
    line-height: 1.22;
    margin-bottom: 1.5px;
    text-align: justify;
  }}
  .skills-row {{
    font-size: 8.8pt;
    line-height: 1.25;
    margin-bottom: 2px;
  }}
  .skills-row strong {{
    font-size: 8.8pt;
  }}
</style>
</head>
<body>

<div class="header">
  <div class="name">{data['name']}</div>
  <div class="contact">{data['contact_line']}</div>
</div>

<div class="section-title">{data['edu_title']}</div>
"""
    for edu in data['education']:
        html += f"""
<div class="item">
  <div class="item-header">
    <span>{edu['institution']}</span>
    <span>{edu['location']}</span>
  </div>
  <div class="item-sub">
    <span>{edu['degree']}</span>
    <span>{edu['dates']}</span>
  </div>
  <ul>
"""
        for b in edu['bullets']:
            html += f"    <li>{b}</li>\n"
        html += "  </ul>\n</div>\n"

    html += f'<div class="section-title">{data["exp_title"]}</div>\n'
    for exp in data['experience']:
        html += f"""
<div class="item">
  <div class="item-header">
    <span>{exp['company']}</span>
    <span>{exp['location']}</span>
  </div>
  <div class="item-sub">
    <span>{exp['role']}</span>
    <span>{exp['dates']}</span>
  </div>
  <ul>
"""
        for b in exp['bullets']:
            html += f"    <li>{b}</li>\n"
        html += "  </ul>\n</div>\n"

    html += f'<div class="section-title">{data["lead_title"]}</div>\n'
    for lead in data['leadership']:
        html += f"""
<div class="item">
  <div class="item-header">
    <span>{lead['company']}</span>
    <span>{lead['location']}</span>
  </div>
  <div class="item-sub">
    <span>{lead['role']}</span>
    <span>{lead['dates']}</span>
  </div>
  <ul>
"""
        for b in lead['bullets']:
            html += f"    <li>{b}</li>\n"
        html += "  </ul>\n</div>\n"

    html += f'<div class="section-title">{data["skills_title"]}</div>\n'
    for s_line in data['skills_lines']:
        html += f'<div class="skills-row"><strong>{s_line["label"]}:</strong> {s_line["value"]}</div>\n'

    html += """
</body>
</html>
"""
    with open(filename, 'w', encoding='utf-8') as f:
        f.write(html)
    print(f"Generated HTML: {filename}")

# Data for Uzbek
data_uz = {
    'name': 'Otabek Haydarov',
    'contact_line': 'Toshkent, Oʻzbekiston  |  +998 90 277 11 04  |  otabek11.haydarov@gmail.com  |  otabekhaydarov.me  |  linkedin.com/in/otabek-haydarov',
    'edu_title': 'TA’LIM',
    'education': [
        {
            'institution': 'PDP University',
            'location': 'Toshkent (Sergeli), Oʻzbekiston',
            'degree': 'Bakalavr: Axborot texnologiyalari (Data Analytics yoʻnalishi)',
            'dates': '2025 – 2029',
            'bullets': [
                'Akademik natija: 1-kurs yakuni boʻyicha GPA: 4.26 / 5.0.',
                'Asosiy kurs ishlari va fanlar: Maʼlumotlar bazalarini boshqarish (SQL), Data Science uchun Python, Maʼlumotlar vizualizatsiyasi va BI tahlil, Ehtimollar nazariyasi va amaliy statistika, Algoritmlar va maʼlumotlar tuzilmalari.'
            ]
        }
    ],
    'exp_title': 'ISH TAJRIBASI VA LOYIHALAR',
    'experience': [
        {
            'company': 'E-Commerce Profitability & Unit Economics Analysis',
            'location': 'Toshkent, Oʻzbekiston',
            'role': 'Data Analyst (PostgreSQL, Python, SQL, Power BI)',
            'dates': '2025',
            'bullets': [
                'PostgreSQL maʼlumotlar bazasida 2,000 dan ortiq real buyurtma tranzaksiyalarini audit qilib, yetkazib berishdagi yoʻqotishlar (margin leakage) 31% dan 13% gacha farq qilishini aniqladim va unit-iqtisodiyotini tahlil qildim.',
                'Turli marketing kanallari boʻyicha investitsiya samaradorligi (ROAS), oʻrtacha buyurtma qiymati (AOV) va mijozning umrboqiy qiymatini (CLV) hisoblovchi murakkab SQL soʻrovlarini (CTEs, Window Functions) ishlab chiqdim.',
                'Rahbariyat uchun sof foyda dinamikasi, mahsulot toifalari rentabelligi va mijozlar xatti-harakatlarini aks ettiruvchi interaktiv Power BI dashboardini yaratdim.'
            ]
        },
        {
            'company': 'U.S. Electric Vehicle Market Share & Infrastructure Analysis',
            'location': 'Toshkent, Oʻzbekiston',
            'role': 'Data Analyst (PostgreSQL, Python, Tableau, SQL)',
            'dates': '2024 – 2025',
            'bullets': [
                '51 ta yurisdiksiya boʻyicha 287 milliondan ortiq roʻyxatga olingan avtotransport vositalariga oid ulkan maʼlumotlar toʻplamini tozalash, normallashtirish va statistik tahlil qilish jarayonini bajardim.',
                'Elektromobillarning (EV) ommalashish surʼatlari, anʼanaviy yoqilgʻi turlaridan oʻtish tendensiyalari va quvvatlash infratuzilmasi talablarini vizual tahlil qilib, Tableau platformasida interaktiv tahliliy panellar yaratdim.',
                'Hududiy quvvatlash stansiyalarini optimal taqsimlash va infratuzilma investitsiyalarini rejalashtirish boʻyicha tavsiyaviy maʼlumotlar modelini shakllantirdim.'
            ]
        },
        {
            'company': 'Natural Language SQL Query Engine',
            'location': 'Toshkent, Oʻzbekiston',
            'role': 'Data & Software Developer (Python, FastAPI, PostgreSQL, LLMs)',
            'dates': '2025',
            'bullets': [
                'Texnik boʻlmagan xodimlar uchun oddiy tildagi soʻrovlarni avtomatik tarzda toʻgʻri SQL kodlariga aylantirib, maʼlumotlar bazasidan tezkor javob oluvchi tizim integratsiyasini qildim.',
                'Maʼlumotlar xavfsizligini taʼminlash uchun SQL-validatsiya va himoya qatlamini joriy etdim, bu esa nojoʻya yoki ruxsat etilmagan amallarning oldini oldi.'
            ]
        }
    ],
    'lead_title': 'ETAKCHILIK TAJRIBASI (LEADERSHIP EXPERIENCE)',
    'leadership': [
        {
            'company': 'AnalystBuilder & Open-Source Community',
            'location': 'Toshkent / Masofaviy',
            'role': 'Data hamjamiyat aʼzosi & Oʻquv guruhi yetakchisi',
            'dates': '2024 – Hozirgacha',
            'bullets': [
                'Alex The Analyst (AnalystBuilder) platformasida maʼlumotlarni tozalash (Data Cleaning), qidiruv tahlili (EDA) va biznes-keyslar boʻyicha amaliy topshiriqlarni muvaffaqiyatli topshirdim.',
                'PDP University talabalari doirasida maʼlumotlar tahlili oʻquv guruhlarida faol qatnashib, tengdoshlarimga murakkab SQL soʻrovlari va Python sintaksisini oʻzlashtirishda koʻmaklashdim.'
            ]
        }
    ],
    'skills_title': 'QIZIQISHLARINGIZ / KO’NIKMALARINGIZ',
    'skills_lines': [
        {
            'label': 'Ko’nikmalar',
            'value': 'Power BI, Tableau, Advanced Excel (Pivot Tables, XLOOKUP, Data Modeling), Python (Pandas, NumPy, Matplotlib), SQL (PostgreSQL, MySQL), JavaScript, TypeScript, React, FastAPI, Git, GitHub, Figma.'
        },
        {
            'label': 'Tillar',
            'value': 'Oʻzbek tili (Ona tili), Tojik tili (Ona tili), Ingliz tili (B2 — Professional ishchi daraja), Rus tili (A1–A2 — Boshlangʻich / Oʻrganilmoqda).'
        },
        {
            'label': 'Qiziqishlar',
            'value': 'Sunʼiy intellekt va LLM vositalari, Maʼlumotlar vizualizatsiyasi dizayni, Shaxmat, Texnologik tahlillar va ochiq kodli loyihalar (Open-source).'
        }
    ]
}

# Data for English
data_en = {
    'name': 'Otabek Haydarov',
    'contact_line': 'Tashkent, Uzbekistan  |  +998 90 277 11 04  |  otabek11.haydarov@gmail.com  |  otabekhaydarov.me  |  linkedin.com/in/otabek-haydarov',
    'edu_title': 'EDUCATION',
    'education': [
        {
            'institution': 'PDP University',
            'location': 'Tashkent, Uzbekistan',
            'degree': 'Bachelor of Science in Information Technology (Data Analytics Track)',
            'dates': '2025 – 2029',
            'bullets': [
                'Academic Standing: First Year GPA: 4.26 / 5.0.',
                'Relevant Coursework: Database Management Systems (SQL), Python for Data Science, Data Visualization & BI, Probability & Applied Statistics, Data Structures and Algorithms.'
            ]
        }
    ],
    'exp_title': 'WORK EXPERIENCE & PROJECTS',
    'experience': [
        {
            'company': 'E-Commerce Profitability & Unit Economics Analysis',
            'location': 'Tashkent, Uzbekistan',
            'role': 'Data Analyst (PostgreSQL, Python, SQL, Power BI)',
            'dates': '2025',
            'bullets': [
                'Audited and transformed a transactional database of 2,000+ customer orders in PostgreSQL to model unit economics and identify hidden margin leakages (pinpointing a 31% to 13% disparity).',
                'Formulated and tuned complex SQL queries using Common Table Expressions (CTEs), Window Functions, and multi-table joins to measure multi-channel marketing ROAS and calculate Customer Lifetime Value (CLV).',
                'Designed an executive Power BI dashboard visualizing gross margin trends, Average Order Value (AOV), and customer cohort behaviors to support commercial strategy.'
            ]
        },
        {
            'company': 'U.S. Electric Vehicle Market Share & Infrastructure Analysis',
            'location': 'Tashkent, Uzbekistan',
            'role': 'Data Analyst (PostgreSQL, Python, Tableau, SQL)',
            'dates': '2024 – 2025',
            'bullets': [
                'Performed end-to-end data cleaning, normalization, and statistical analysis on a dataset of 287M+ registered vehicles across 51 jurisdictions.',
                'Uncovered adoption velocity trends, alternative fuel penetration, and charging infrastructure bottlenecks using Python and PostgreSQL.',
                'Built dynamic, interactive Tableau visualizations delivering actionable geographic insights for charging station capacity planning and policy analysis.'
            ]
        },
        {
            'company': 'Natural Language SQL Query Engine',
            'location': 'Tashkent, Uzbekistan',
            'role': 'Data & Software Developer (Python, FastAPI, PostgreSQL, LLMs)',
            'dates': '2025',
            'bullets': [
                'Engineered an analytics tool enabling non-technical stakeholders to input business questions in natural English and obtain validated SQL query outputs in real time.',
                'Implemented a schema-aware validation and safety guard layer to block unintended data mutations and ensure query syntactic accuracy.'
            ]
        }
    ],
    'lead_title': 'LEADERSHIP EXPERIENCE',
    'leadership': [
        {
            'company': 'AnalystBuilder & Open-Source Community',
            'location': 'Tashkent / Remote',
            'role': 'Data Analytics Trainee & Peer Tutor',
            'dates': '2024 – Present',
            'bullets': [
                'Completed rigorous practical business analytics tracks on AnalystBuilder (founded by Alex The Analyst), mastering exploratory data analysis (EDA), data cleaning, and relational database modeling.',
                'Facilitated study groups among university peers at PDP University, assisting fellow students with SQL query optimization and Python data manipulation techniques.'
            ]
        }
    ],
    'skills_title': 'SKILLS & INTERESTS',
    'skills_lines': [
        {
            'label': 'Skills',
            'value': 'Power BI, Tableau, Advanced Excel (Pivot Tables, XLOOKUP, Nested Formulas, Data Modeling), Python (Pandas, NumPy, Matplotlib), SQL (PostgreSQL, MySQL), JavaScript, TypeScript, React, FastAPI, Git, GitHub, Figma.'
        },
        {
            'label': 'Languages',
            'value': 'Uzbek (Native), Tajik (Native), English (B2 – Professional Working Proficiency), Russian (A1–A2 – Elementary / Actively Learning).'
        },
        {
            'label': 'Interests',
            'value': 'AI & Large Language Models, Data Visualization Design, Chess, Open-Source Technologies, Tech Blogging.'
        }
    ]
}

if __name__ == '__main__':
    base_dir = os.path.abspath(os.path.dirname(__file__))
    
    # Generate DOCX
    docx_uz = os.path.join(base_dir, 'Resume_Otabek_Haydarov_UZ.docx')
    docx_en = os.path.join(base_dir, 'Resume_Otabek_Haydarov_EN.docx')
    create_resume_docx(data_uz, docx_uz)
    create_resume_docx(data_en, docx_en)
    
    # Generate HTML
    html_uz = os.path.join(base_dir, 'Resume_Otabek_Haydarov_UZ.html')
    html_en = os.path.join(base_dir, 'Resume_Otabek_Haydarov_EN.html')
    create_resume_html(data_uz, html_uz)
    create_resume_html(data_en, html_en)
