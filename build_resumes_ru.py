import os
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_resume_docx(data, filename):
    doc = docx.Document()
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(10)
    font.color.rgb = RGBColor(0, 0, 0)

    def add_bottom_border(paragraph):
        pPr = paragraph._p.get_or_add_pPr()
        pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                         r'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="000000"/>'
                         r'</w:pBdr>')
        pPr.append(pBdr)

    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.name = 'Times New Roman'
        add_bottom_border(p)
        return p

    def add_entry_header(left_title, right_info, left_sub="", right_sub=""):
        table = doc.add_table(rows=0, cols=2)
        table.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        
        row1 = table.add_row()
        cell_l1 = row1.cells[0]
        cell_r1 = row1.cells[1]
        cell_l1.width = Inches(5.6)
        cell_r1.width = Inches(1.9)
        
        p_l1 = cell_l1.paragraphs[0]
        p_l1.paragraph_format.space_before = Pt(3)
        p_l1.paragraph_format.space_after = Pt(0)
        run_l1 = p_l1.add_run(left_title)
        run_l1.bold = True
        run_l1.font.size = Pt(10)
        run_l1.font.name = 'Times New Roman'
        
        p_r1 = cell_r1.paragraphs[0]
        p_r1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_r1.paragraph_format.space_before = Pt(3)
        p_r1.paragraph_format.space_after = Pt(0)
        run_r1 = p_r1.add_run(right_info)
        run_r1.bold = True
        run_r1.font.size = Pt(10)
        run_r1.font.name = 'Times New Roman'
        
        if left_sub or right_sub:
            row2 = table.add_row()
            cell_l2 = row2.cells[0]
            cell_r2 = row2.cells[1]
            cell_l2.width = Inches(5.6)
            cell_r2.width = Inches(1.9)
            
            p_l2 = cell_l2.paragraphs[0]
            p_l2.paragraph_format.space_before = Pt(0)
            p_l2.paragraph_format.space_after = Pt(1)
            run_l2 = p_l2.add_run(left_sub)
            run_l2.italic = True
            run_l2.font.size = Pt(9.5)
            run_l2.font.name = 'Times New Roman'
            
            p_r2 = cell_r2.paragraphs[0]
            p_r2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_r2.paragraph_format.space_before = Pt(0)
            p_r2.paragraph_format.space_after = Pt(1)
            run_r2 = p_r2.add_run(right_sub)
            run_r2.italic = True
            run_r2.font.size = Pt(9.5)
            run_r2.font.name = 'Times New Roman'

    def add_bullet(text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.left_indent = Inches(0.2)
        run = p.add_run(text)
        run.font.size = Pt(9)
        run.font.name = 'Times New Roman'

    # Name & Contact
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(1)
    run_name = p_name.add_run(data['name'].upper())
    run_name.bold = True
    run_name.font.size = Pt(16)
    run_name.font.name = 'Times New Roman'

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.space_after = Pt(4)
    run_contact = p_contact.add_run(data['contact_line'])
    run_contact.font.size = Pt(9)
    run_contact.font.name = 'Times New Roman'

    # Education
    add_section_header(data['edu_title'])
    for edu in data['education']:
        add_entry_header(edu['institution'], edu['location'], edu['degree'], edu['dates'])
        for bullet in edu['bullets']:
            add_bullet(bullet)

    # Experience
    add_section_header(data['exp_title'])
    for exp in data['experience']:
        add_entry_header(exp['company'], exp['location'], exp['role'], exp['dates'])
        for bullet in exp['bullets']:
            add_bullet(bullet)

    # Leadership
    add_section_header(data['lead_title'])
    for lead in data['leadership']:
        add_entry_header(lead['company'], lead['location'], lead['role'], lead['dates'])
        for bullet in lead['bullets']:
            add_bullet(bullet)

    # Skills
    add_section_header(data['skills_title'])
    for s_line in data['skills_lines']:
        p_skill = doc.add_paragraph()
        p_skill.paragraph_format.space_before = Pt(0)
        p_skill.paragraph_format.space_after = Pt(1)
        p_skill.paragraph_format.line_spacing = 1.05
        run_bold = p_skill.add_run(s_line['label'] + ": ")
        run_bold.bold = True
        run_bold.font.size = Pt(9)
        run_bold.font.name = 'Times New Roman'
        run_val = p_skill.add_run(s_line['value'])
        run_val.font.size = Pt(9)
        run_val.font.name = 'Times New Roman'

    doc.save(filename)
    print(f"Generated DOCX: {filename}")

def create_resume_html(data, filename):
    html = f"""<!DOCTYPE html>
<html lang="ru">
<head>
<meta charset="UTF-8">
<style>
  @page {{
    size: letter;
    margin: 0.45in 0.5in;
  }}
  * {{
    box-sizing: border-box;
    margin: 0;
    padding: 0;
  }}
  body {{
    font-family: 'Times New Roman', Times, serif;
    font-size: 9.5pt;
    line-height: 1.25;
    color: #111;
    background-color: #fff;
  }}
  .header {{
    text-align: center;
    margin-bottom: 6px;
  }}
  .name {{
    font-size: 16pt;
    font-weight: bold;
    letter-spacing: 0.5px;
    margin-bottom: 2px;
    text-transform: uppercase;
  }}
  .contact {{
    font-size: 8.8pt;
    color: #222;
  }}
  .contact a {{
    color: inherit;
    text-decoration: none;
  }}
  .section-title {{
    font-size: 10.5pt;
    font-weight: bold;
    text-transform: uppercase;
    border-bottom: 1px solid #111;
    margin-top: 7px;
    margin-bottom: 3px;
    padding-bottom: 1px;
    letter-spacing: 0.4px;
  }}
  .item {{
    margin-bottom: 3px;
  }}
  .item-header {{
    display: flex;
    justify-content: space-between;
    font-size: 9.5pt;
    font-weight: bold;
  }}
  .item-sub {{
    display: flex;
    justify-content: space-between;
    font-size: 9pt;
    font-style: italic;
    color: #222;
    margin-bottom: 1px;
  }}
  ul {{
    margin: 1px 0 2px 18px;
    padding-left: 0;
  }}
  li {{
    font-size: 8.8pt;
    line-height: 1.22;
    margin-bottom: 1.5px;
    text-align: justify;
  }}
  .skills-row {{
    font-size: 8.8pt;
    line-height: 1.25;
    margin-bottom: 2px;
  }}
  .skills-row strong {{
    font-size: 8.8pt;
  }}
</style>
</head>
<body>

<div class="header">
  <div class="name">{data['name']}</div>
  <div class="contact">{data['contact_line']}</div>
</div>

<div class="section-title">{data['edu_title']}</div>
"""
    for edu in data['education']:
        html += f"""
<div class="item">
  <div class="item-header">
    <span>{edu['institution']}</span>
    <span>{edu['location']}</span>
  </div>
  <div class="item-sub">
    <span>{edu['degree']}</span>
    <span>{edu['dates']}</span>
  </div>
  <ul>
"""
        for b in edu['bullets']:
            html += f"    <li>{b}</li>\n"
        html += "  </ul>\n</div>\n"

    html += f'<div class="section-title">{data["exp_title"]}</div>\n'
    for exp in data['experience']:
        html += f"""
<div class="item">
  <div class="item-header">
    <span>{exp['company']}</span>
    <span>{exp['location']}</span>
  </div>
  <div class="item-sub">
    <span>{exp['role']}</span>
    <span>{exp['dates']}</span>
  </div>
  <ul>
"""
        for b in exp['bullets']:
            html += f"    <li>{b}</li>\n"
        html += "  </ul>\n</div>\n"

    html += f'<div class="section-title">{data["lead_title"]}</div>\n'
    for lead in data['leadership']:
        html += f"""
<div class="item">
  <div class="item-header">
    <span>{lead['company']}</span>
    <span>{lead['location']}</span>
  </div>
  <div class="item-sub">
    <span>{lead['role']}</span>
    <span>{lead['dates']}</span>
  </div>
  <ul>
"""
        for b in lead['bullets']:
            html += f"    <li>{b}</li>\n"
        html += "  </ul>\n</div>\n"

    html += f'<div class="section-title">{data["skills_title"]}</div>\n'
    for s_line in data['skills_lines']:
        html += f'<div class="skills-row"><strong>{s_line["label"]}:</strong> {s_line["value"]}</div>\n'

    html += """
</body>
</html>
"""
    with open(filename, 'w', encoding='utf-8') as f:
        f.write(html)
    print(f"Generated HTML: {filename}")

data_ru = {
    'name': 'Отабек Хайдаров',
    'contact_line': 'Ташкент, Узбекистан  |  +998 90 277 11 04  |  otabek11.haydarov@gmail.com  |  otabekhaydarov.me  |  linkedin.com/in/otabek-haydarov',
    'edu_title': 'ОБРАЗОВАНИЕ',
    'education': [
        {
            'institution': 'PDP University',
            'location': 'Ташкент (Сергели), Узбекистан',
            'degree': 'Бакалавриат: Информационные технологии (Направление: Аналитика данных / Data Analytics)',
            'dates': '2025 – 2029',
            'bullets': [
                'Академическая успеваемость: GPA 4.26 / 5.0 (по итогам 1 курса).',
                'Профильные дисциплины: Системы управления базами данных (SQL), Python для Data Science, Визуализация данных и BI-аналитика, Теория вероятностей и прикладная статистика, Алгоритмы и структуры данных.'
            ]
        }
    ],
    'exp_title': 'ОПЫТ РАБОТЫ И ПРОЕКТЫ',
    'experience': [
        {
            'company': 'E-Commerce Profitability & Unit Economics Analysis',
            'location': 'Ташкент, Узбекистан',
            'role': 'Data Analyst (PostgreSQL, Python, SQL, Power BI)',
            'dates': '2025',
            'bullets': [
                'Провел детальный аудит транзакционной базы данных (2,000+ заказов) в PostgreSQL для оценки unit-экономики и выявления скрытых утечек маржинальности (margin leakage от 31% до 13%).',
                'Спроектировал и оптимизировал сложные аналитические SQL-запросы (CTE, оконные функции, многотабличные JOIN) для расчета маркетингового ROAS, среднего чека (AOV) и LTV клиента.',
                'Разработал интерактивный дашборд в Power BI для руководства, визуализирующий динамику чистой прибыли, рентабельность товарных категорий и поведение когорт покупателей.'
            ]
        },
        {
            'company': 'U.S. Electric Vehicle Market Share & Infrastructure Analysis',
            'location': 'Ташкент, Узбекистан',
            'role': 'Data Analyst (PostgreSQL, Python, Tableau, SQL)',
            'dates': '2024 – 2025',
            'bullets': [
                'Осуществил полный цикл очистки, нормализации и статистического анализа массива данных по 287+ млн зарегистрированных транспортных средств в 51 юрисдикции.',
                'Выявил динамику проникновения электромобилей (EV), тренды перехода с традиционного топлива и дефицит зарядной инфраструктуры с помощью Python и PostgreSQL.',
                'Построил динамические дашборды в Tableau, предоставив стейкхолдерам пространственную аналитику для регионального планирования инвестиций в электрозарядные станции.'
            ]
        },
        {
            'company': 'Natural Language SQL Query Engine',
            'location': 'Ташкент, Узбекистан',
            'role': 'Data & Software Developer (Python, FastAPI, PostgreSQL, LLMs)',
            'dates': '2025',
            'bullets': [
                'Разработал инструмент, позволяющий нетехническим сотрудникам формулировать бизнес-запросы на естественном языке и мгновенно получать валидные SQL-выборки из БД.',
                'Реализовал защитный слой валидации структуры запросов (guard layer) на базе FastAPI, предотвращающий выполнение некорректных и несанкционированных операций.'
            ]
        }
    ],
    'lead_title': 'ЛИДЕРСКИЙ ОПЫТ (LEADERSHIP EXPERIENCE)',
    'leadership': [
        {
            'company': 'AnalystBuilder & Open-Source Community',
            'location': 'Ташкент / Удаленно',
            'role': 'Участник Data-сообщества & Ментор учебных групп',
            'dates': '2024 – настоящее время',
            'bullets': [
                'Успешно выполнил серию прикладных бизнес-кейсов на платформе AnalystBuilder (автор: Alex The Analyst) по очистке данных, разведочному анализу (EDA) и реляционному моделированию.',
                'Курировал академические группы студентов в PDP University, консультируя коллег по синтаксису и оптимизации сложных SQL-запросов, а также библиотекам Pandas и NumPy.'
            ]
        }
    ],
    'skills_title': 'НАВЫКИ И ИНТЕРЕСЫ',
    'skills_lines': [
        {
            'label': 'Навыки',
            'value': 'Power BI, Tableau, Продвинутый Excel (Сводные таблицы, XLOOKUP, Data Modeling), Python (Pandas, NumPy, Matplotlib), SQL (PostgreSQL, MySQL), JavaScript, TypeScript, React, FastAPI, Git, GitHub, Figma.'
        },
        {
            'label': 'Языки',
            'value': 'Узбекский (Родной), Таджикский (Родной), Английский (B2 — Профессиональный рабочий уровень), Русский (A1–A2 — Базовый / Активно изучается).'
        },
        {
            'label': 'Интересы',
            'value': 'Искусственный интеллект и LLM, дизайн визуализации данных, шахматы, Open-source разработка, технологическая аналитика.'
        }
    ]
}

if __name__ == '__main__':
    base_dir = os.path.abspath(os.path.dirname(__file__))
    
    docx_ru = os.path.join(base_dir, 'Resume_Otabek_Haydarov_RU.docx')
    html_ru = os.path.join(base_dir, 'Resume_Otabek_Haydarov_RU.html')
    
    create_resume_docx(data_ru, docx_ru)
    create_resume_html(data_ru, html_ru)
